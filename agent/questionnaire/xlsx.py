import io

import openpyxl


def parse_xlsx(file_bytes, question_col="A", header_rows=0):
    """Extract questions from a column of an .xlsx questionnaire.

    Returns [{"row", "question"}] for every non-blank cell below the header rows,
    preserving the original row number so answers can be written back in place.
    """
    ws = openpyxl.load_workbook(io.BytesIO(file_bytes)).active
    out = []
    for cell in ws[question_col]:
        if cell.row <= header_rows:
            continue
        if cell.value is None or str(cell.value).strip() == "":
            continue
        out.append({"row": cell.row, "question": cell.value})
    return out


def export_xlsx(file_bytes, results, answer_col="B"):
    """Write each result's answer back into the original questionnaire at its row.

    The original workbook (questions, layout, other columns) is preserved; only the
    answer column is filled, so the deliverable matches the format the buyer sent.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = wb.active
    for result in results:
        ws[f"{answer_col}{result['row']}"] = result["answer"]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
