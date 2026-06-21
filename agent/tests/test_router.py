"""TDD for the FastAPI router.

Uses TestClient + FastAPI dependency-override so the full HTTP path
(billing gate -> parse -> answer -> export) is exercised with fakes injected only
at the I/O boundary (Gemini, evidence, log, account, usage, checkout). No SDK, no network.
"""

import io
import json

import openpyxl
from fastapi.testclient import TestClient

import main
from main import (
    app,
    get_account,
    get_checkout_link,
    get_evidence_provider,
    get_generate,
    get_logger,
    get_mark_paid,
    get_usage_recorder,
    get_webhook_parser,
)


def _stub_providers():
    """Safe defaults for every injected dependency; an entitled (free, unused) account."""
    app.dependency_overrides[get_generate] = lambda: (lambda _p: "{}")
    app.dependency_overrides[get_evidence_provider] = lambda: (lambda _c: "")
    app.dependency_overrides[get_logger] = lambda: (lambda _r: None)
    app.dependency_overrides[get_account] = lambda: (
        lambda _c: {"questionnaires_used": 0, "is_paid": False}
    )
    app.dependency_overrides[get_usage_recorder] = lambda: (lambda _c: None)
    app.dependency_overrides[get_checkout_link] = lambda: (lambda _c: "https://stripe.test/checkout")


def _xlsx(questions):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, q in enumerate(questions, start=1):
        ws[f"A{i}"] = q
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, book, customer_id="acme"):
    return client.post(
        "/questionnaires/answer",
        data={"customer_id": customer_id},
        files={
            "file": (
                "q.xlsx",
                book,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def test_endpoint_passes_customer_id_to_evidence_provider():
    seen = {}

    def evidence(customer_id):
        seen["customer_id"] = customer_id
        return "evidence"

    _stub_providers()
    app.dependency_overrides[get_evidence_provider] = lambda: evidence
    try:
        resp = _post(TestClient(app), _xlsx(["Do you have MFA?"]), customer_id="acme-42")
    finally:
        app.dependency_overrides.clear()

    assert resp.status_code == 200
    assert seen["customer_id"] == "acme-42"


def test_endpoint_returns_completed_xlsx_with_counts_and_logs_decisions():
    logged = []

    def fake_generate(prompt):
        if "encrypt" in prompt:
            return json.dumps(
                {"can_answer": True, "confidence": 0.95, "answer": "Yes, AES-256", "citation": "SOC2"}
            )
        return json.dumps({"can_answer": False, "confidence": 0.1})

    _stub_providers()
    app.dependency_overrides[get_generate] = lambda: fake_generate
    app.dependency_overrides[get_evidence_provider] = lambda: (lambda _c: "<evidence>")
    app.dependency_overrides[get_logger] = lambda: logged.append
    try:
        resp = _post(TestClient(app), _xlsx(["Do you encrypt data at rest?", "Do you do annual pentests?"]))
    finally:
        app.dependency_overrides.clear()

    assert resp.status_code == 200
    assert resp.headers["x-total"] == "2"
    assert resp.headers["x-answered"] == "1"
    assert resp.headers["x-deferred"] == "1"

    ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
    assert ws["B1"].value == "Yes, AES-256"
    assert ws["B2"].value == "NEEDS YOUR INPUT"
    assert len(logged) == 2


def test_logged_records_are_enriched_for_bigquery():
    logged = []
    _stub_providers()
    app.dependency_overrides[get_generate] = lambda: (
        lambda _p: json.dumps({"can_answer": False, "confidence": 0.0})
    )
    app.dependency_overrides[get_evidence_provider] = lambda: (lambda _c: "<evidence>")
    app.dependency_overrides[get_logger] = lambda: logged.append
    try:
        _post(TestClient(app), _xlsx(["Do you have MFA?", "Do you log access?"]), customer_id="acme-9")
    finally:
        app.dependency_overrides.clear()

    assert all(r["customer_id"] == "acme-9" for r in logged)
    assert all(r.get("run_id") for r in logged)
    assert len({r["run_id"] for r in logged}) == 1
    assert all(r.get("created_at") for r in logged)


def test_unpaid_customer_over_quota_is_blocked_with_checkout_link():
    _stub_providers()
    app.dependency_overrides[get_account] = lambda: (
        lambda _c: {"questionnaires_used": 1, "is_paid": False}
    )
    app.dependency_overrides[get_checkout_link] = lambda: (lambda _c: "https://stripe.test/checkout/abc")
    try:
        resp = _post(TestClient(app), _xlsx(["Do you have MFA?"]))
    finally:
        app.dependency_overrides.clear()

    assert resp.status_code == 402
    assert resp.json()["detail"]["checkout_url"] == "https://stripe.test/checkout/abc"


def test_paid_customer_over_quota_is_answered():
    _stub_providers()
    app.dependency_overrides[get_account] = lambda: (
        lambda _c: {"questionnaires_used": 99, "is_paid": True}
    )
    try:
        resp = _post(TestClient(app), _xlsx(["Do you have MFA?"]))
    finally:
        app.dependency_overrides.clear()

    assert resp.status_code == 200


def test_usage_recorded_after_answering_but_not_when_blocked():
    answered = []
    _stub_providers()
    app.dependency_overrides[get_usage_recorder] = lambda: answered.append
    try:
        _post(TestClient(app), _xlsx(["Do you have MFA?"]), customer_id="cust-A")
    finally:
        app.dependency_overrides.clear()
    assert answered == ["cust-A"]

    blocked = []
    _stub_providers()
    app.dependency_overrides[get_account] = lambda: (
        lambda _c: {"questionnaires_used": 1, "is_paid": False}
    )
    app.dependency_overrides[get_usage_recorder] = lambda: blocked.append
    try:
        _post(TestClient(app), _xlsx(["Do you have MFA?"]), customer_id="cust-B")
    finally:
        app.dependency_overrides.clear()
    assert blocked == []  # a blocked (402) request must not consume usage


def test_checkout_returns_a_subscription_url():
    app.dependency_overrides[get_checkout_link] = lambda: (lambda c: f"https://stripe.test/{c}")
    try:
        resp = TestClient(app).post("/billing/checkout", data={"customer_id": "acme"})
    finally:
        app.dependency_overrides.clear()
    assert resp.status_code == 200
    assert resp.json()["checkout_url"] == "https://stripe.test/acme"


def test_webhook_marks_customer_paid_on_subscription_event():
    marked = []
    app.dependency_overrides[get_webhook_parser] = lambda: (lambda _payload, _sig: "cust-paid")
    app.dependency_overrides[get_mark_paid] = lambda: marked.append
    try:
        resp = TestClient(app).post("/stripe/webhook", content=b"{}", headers={"stripe-signature": "x"})
    finally:
        app.dependency_overrides.clear()
    assert resp.status_code == 200
    assert marked == ["cust-paid"]


def test_webhook_ignores_unrelated_events():
    marked = []
    app.dependency_overrides[get_webhook_parser] = lambda: (lambda _payload, _sig: None)
    app.dependency_overrides[get_mark_paid] = lambda: marked.append
    try:
        resp = TestClient(app).post("/stripe/webhook", content=b"{}", headers={"stripe-signature": "x"})
    finally:
        app.dependency_overrides.clear()
    assert resp.status_code == 200
    assert marked == []


def test_rejects_oversized_upload(monkeypatch):
    monkeypatch.setattr(main, "MAX_UPLOAD_BYTES", 100, raising=False)
    _stub_providers()
    try:
        resp = _post(TestClient(app), _xlsx(["Do you have MFA?"]))  # a real .xlsx is > 100 bytes
    finally:
        app.dependency_overrides.clear()
    assert resp.status_code == 413


def test_rejects_non_xlsx_upload():
    _stub_providers()
    try:
        resp = TestClient(app).post(
            "/questionnaires/answer",
            data={"customer_id": "acme"},
            files={"file": ("q.xlsx", b"this is not a spreadsheet",
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    finally:
        app.dependency_overrides.clear()
    assert resp.status_code == 400
