"""TDD for the XLSX questionnaire round-trip (T-03 parse + T-08 export).

These tests build real .xlsx workbooks in memory with openpyxl — no mocking — so
the parse/export behavior is verified against actual spreadsheet bytes.
"""

import io
import json

import openpyxl

from questionnaire.batch import answer_questionnaire
from questionnaire.decision import NEEDS_INPUT
from questionnaire.xlsx import parse_xlsx, export_xlsx


def _build_xlsx(column_a):
    """Build .xlsx bytes with the given values down column A (None = leave blank)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, value in enumerate(column_a, start=1):
        if value is not None:
            ws[f"A{i}"] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_extracts_questions_with_row_numbers():
    book = _build_xlsx(["Do you encrypt data at rest?", "Do you have MFA?"])
    assert parse_xlsx(book) == [
        {"row": 1, "question": "Do you encrypt data at rest?"},
        {"row": 2, "question": "Do you have MFA?"},
    ]


def test_parse_skips_blank_cells():
    book = _build_xlsx(["Q1?", None, "Q3?"])
    parsed = parse_xlsx(book)
    assert [p["row"] for p in parsed] == [1, 3]
    assert [p["question"] for p in parsed] == ["Q1?", "Q3?"]


def test_parse_skips_header_rows():
    book = _build_xlsx(["Question", "Do you have MFA?"])
    assert parse_xlsx(book, header_rows=1) == [{"row": 2, "question": "Do you have MFA?"}]


def test_round_trip_writes_answers_next_to_their_questions():
    book = _build_xlsx(["Do you encrypt data at rest?", "Do you have MFA?"])
    parsed = parse_xlsx(book)
    results = [
        {"row": p["row"], "question": p["question"], "answer": f"A: {p['question']}"}
        for p in parsed
    ]

    out = export_xlsx(book, results)

    ws = openpyxl.load_workbook(io.BytesIO(out)).active
    assert ws["A1"].value == "Do you encrypt data at rest?"  # original question preserved
    assert ws["B1"].value == "A: Do you encrypt data at rest?"
    assert ws["A2"].value == "Do you have MFA?"
    assert ws["B2"].value == "A: Do you have MFA?"


def test_full_chain_parse_answer_export_blanks_deferred_rows():
    book = _build_xlsx(["Do you encrypt data at rest?", "Do you do annual pentests?"])
    parsed = parse_xlsx(book)

    def generate(prompt):
        if "encrypt" in prompt:
            return json.dumps(
                {"can_answer": True, "confidence": 0.95, "answer": "Yes, AES-256", "citation": "SOC2"}
            )
        return json.dumps({"can_answer": False, "confidence": 0.1})

    run = answer_questionnaire(parsed, "<evidence>", generate)
    out = export_xlsx(book, run["results"])

    ws = openpyxl.load_workbook(io.BytesIO(out)).active
    assert ws["B1"].value == "Yes, AES-256"
    assert ws["B2"].value == NEEDS_INPUT  # deferred answer exported as a clear human prompt
